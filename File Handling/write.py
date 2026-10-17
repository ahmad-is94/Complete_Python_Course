# import openpyxl module
import openpyxl

# we are calling a Workbook() function of openpyxl
wb = openpyxl.Workbook()

sheet = wb.active
c1 = sheet.cell(row=4, column=1)

# writing values to cells
c1.value = "A"

c2 = sheet.cell(row=1, column=2)
c2.value = "S"
c3 = sheet['C1']
c3.value = "CONTENT DEVELOPER"

# A2 means column = 1 & row = 2.
c4 = sheet['A2']
c4.value = "V"

# B2 means column = 2 & row = 2.
c5 = sheet['B2']
c5.value = "S"
# printing in 2nd row and 3rd column
c6 = sheet['C2']
c6.value = "CONTENT REVIEWER"

# We are saving the file using the save() method in a specific location we want
wb.save(r"E:\Python_Course\Complete_Python_Course\write.xlsx")